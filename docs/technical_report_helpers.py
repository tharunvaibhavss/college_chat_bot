"""
Generator for MCA Academic Assistant Technical Project Report (.docx and .pdf).
Modeled directly after the high-density technical report structure and aesthetic
of the reference 'Hospital Patient Scheduling Report'.
"""
import os
import sys
import docx
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
from docx.oxml import parse_xml, OxmlElement
from docx.oxml.ns import nsdecls, qn
import win32com.client


def set_cell_background(cell, fill_hex):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{fill_hex}"/>')
    tcPr.append(shd)


def set_cell_margins(cell, top=80, bottom=80, left=120, right=120):
    tcPr = cell._tc.get_or_add_tcPr()
    tcMar = parse_xml(
        f'<w:tcMar {nsdecls("w")}>'
        f'<w:top w:w="{top}" w:type="dxa"/>'
        f'<w:bottom w:w="{bottom}" w:type="dxa"/>'
        f'<w:left w:w="{left}" w:type="dxa"/>'
        f'<w:right w:w="{right}" w:type="dxa"/>'
        f'</w:tcMar>'
    )
    tcPr.append(tcMar)


def add_hyperlink(paragraph, url, text, color="004B87", underline=True, bold=False, font_size=None):
    part = paragraph.part
    r_id = part.relate_to(url, docx.opc.constants.RELATIONSHIP_TYPE.HYPERLINK, is_external=True)

    hyperlink = parse_xml(
        f'<w:hyperlink {nsdecls("w")} r:id="{r_id}" '
        f'xmlns:r="http://schemas.openxmlformats.org/officeDocument/2006/relationships"/>'
    )
    new_run = parse_xml(f'<w:r {nsdecls("w")}/>')
    rPr = parse_xml(f'<w:rPr {nsdecls("w")}/>')

    if color:
        c = parse_xml(f'<w:color {nsdecls("w")} w:val="{color}"/>')
        rPr.append(c)
    if underline:
        u = parse_xml(f'<w:u {nsdecls("w")} w:val="single"/>')
        rPr.append(u)
    if bold:
        b = parse_xml(f'<w:b {nsdecls("w")}/>')
        rPr.append(b)
    if font_size:
        sz = parse_xml(f'<w:sz {nsdecls("w")} w:val="{int(font_size * 2)}"/>')
        rPr.append(sz)

    rFonts = parse_xml(f'<w:rFonts {nsdecls("w")} w:ascii="Times New Roman" w:hAnsi="Times New Roman"/>')
    rPr.append(rFonts)

    new_run.append(rPr)
    text_node = parse_xml(f'<w:t {nsdecls("w")}>{text}</w:t>')
    new_run.append(text_node)
    hyperlink.append(new_run)
    paragraph._p.append(hyperlink)
    return hyperlink


def style_table(table, col_widths=None, header_bg="0F172A", alt_bg="F8FAFC", border_c="CBD5E1"):
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    tblPr = table._tbl.tblPr
    tblBorders = parse_xml(
        f'<w:tblBorders {nsdecls("w")}>'
        f'<w:top w:val="single" w:sz="6" w:space="0" w:color="{border_c}"/>'
        f'<w:bottom w:val="single" w:sz="6" w:space="0" w:color="{border_c}"/>'
        f'<w:left w:val="single" w:sz="6" w:space="0" w:color="{border_c}"/>'
        f'<w:right w:val="single" w:sz="6" w:space="0" w:color="{border_c}"/>'
        f'<w:insideH w:val="single" w:sz="4" w:space="0" w:color="{border_c}"/>'
        f'<w:insideV w:val="single" w:sz="4" w:space="0" w:color="{border_c}"/>'
        f'</w:tblBorders>'
    )
    tblPr.append(tblBorders)

    for row_idx, row in enumerate(table.rows):
        trPr = row._tr.get_or_add_trPr()
        trPr.append(parse_xml(f'<w:cantSplit {nsdecls("w")}/>'))
        if row_idx == 0:
            trPr.append(parse_xml(f'<w:tblHeader {nsdecls("w")}/>'))

        for col_idx, cell in enumerate(row.cells):
            cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
            set_cell_margins(cell, top=60, bottom=60, left=100, right=100)
            if col_widths and col_idx < len(col_widths):
                cell.width = col_widths[col_idx]

            if row_idx == 0:
                set_cell_background(cell, header_bg)
                for p in cell.paragraphs:
                    p.paragraph_format.space_before = Pt(2)
                    p.paragraph_format.space_after = Pt(2)
                    p.paragraph_format.line_spacing = 1.0
                    for r in p.runs:
                        r.font.bold = True
                        r.font.color.rgb = RGBColor(255, 255, 255)
                        r.font.name = "Times New Roman"
                        r.font.size = Pt(9)
            else:
                if row_idx % 2 == 1 and alt_bg:
                    set_cell_background(cell, alt_bg)
                for p in cell.paragraphs:
                    p.paragraph_format.space_before = Pt(2)
                    p.paragraph_format.space_after = Pt(2)
                    p.paragraph_format.line_spacing = 1.05
                    for r in p.runs:
                        r.font.name = "Times New Roman"
                        r.font.size = Pt(8.5)
                        r.font.color.rgb = RGBColor(15, 23, 42)


def add_image_box(doc, image_path, fig_no, title, explanation, width=Inches(6.2)):
    if os.path.exists(image_path):
        p_img = doc.add_paragraph()
        p_img.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p_img.paragraph_format.space_before = Pt(6)
        p_img.paragraph_format.space_after = Pt(3)
        p_img.paragraph_format.keep_with_next = True
        p_img.add_run().add_picture(image_path, width=width)

    p_cap = doc.add_paragraph()
    p_cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_cap.paragraph_format.space_before = Pt(2)
    p_cap.paragraph_format.space_after = Pt(3)
    p_cap.paragraph_format.keep_with_next = True
    r = p_cap.add_run(f"Figure {fig_no}: {title}")
    r.font.bold = True
    r.font.size = Pt(9.5)
    r.font.color.rgb = RGBColor(15, 23, 42)

    p_exp = doc.add_paragraph()
    p_exp.paragraph_format.space_before = Pt(0)
    p_exp.paragraph_format.space_after = Pt(10)
    r_exp_lbl = p_exp.add_run("Explanation: ")
    r_exp_lbl.font.italic = True
    r_exp_lbl.font.size = Pt(9)
    r_exp_lbl.font.color.rgb = RGBColor(71, 85, 105)
    r_exp_txt = p_exp.add_run(explanation)
    r_exp_txt.font.size = Pt(9)
    r_exp_txt.font.color.rgb = RGBColor(51, 65, 85)


def add_code_block(doc, text):
    tbl = doc.add_table(rows=1, cols=1)
    tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    cell = tbl.cell(0, 0)
    cell.width = Inches(6.4)
    set_cell_background(cell, "F8FAFC")
    set_cell_margins(cell, top=100, bottom=100, left=140, right=140)

    tcPr = cell._tc.get_or_add_tcPr()
    tcBorders = parse_xml(
        f'<w:tcBorders {nsdecls("w")}>'
        f'<w:top w:val="single" w:sz="4" w:space="0" w:color="CBD5E1"/>'
        f'<w:bottom w:val="single" w:sz="4" w:space="0" w:color="CBD5E1"/>'
        f'<w:left w:val="single" w:sz="18" w:space="0" w:color="0284C7"/>'
        f'<w:right w:val="single" w:sz="4" w:space="0" w:color="CBD5E1"/>'
        f'</w:tcBorders>'
    )
    tcPr.append(tcBorders)

    p = cell.paragraphs[0]
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(0)
    p.paragraph_format.line_spacing = 1.05
    r = p.add_run(text)
    r.font.name = "Consolas"
    r.font.size = Pt(8.5)
    r.font.color.rgb = RGBColor(15, 23, 42)
    doc.add_paragraph().paragraph_format.space_after = Pt(4)


def add_callout_box(doc, title, body_text):
    tbl = doc.add_table(rows=1, cols=1)
    tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    cell = tbl.cell(0, 0)
    cell.width = Inches(6.4)
    set_cell_background(cell, "EFF6FF")
    set_cell_margins(cell, top=120, bottom=120, left=160, right=160)

    tcPr = cell._tc.get_or_add_tcPr()
    tcBorders = parse_xml(
        f'<w:tcBorders {nsdecls("w")}>'
        f'<w:top w:val="single" w:sz="6" w:space="0" w:color="93C5FD"/>'
        f'<w:bottom w:val="single" w:sz="6" w:space="0" w:color="93C5FD"/>'
        f'<w:left w:val="single" w:sz="6" w:space="0" w:color="93C5FD"/>'
        f'<w:right w:val="single" w:sz="6" w:space="0" w:color="93C5FD"/>'
        f'</w:tcBorders>'
    )
    tcPr.append(tcBorders)

    p = cell.paragraphs[0]
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(0)
    p.paragraph_format.line_spacing = 1.15
    r1 = p.add_run(title + " ")
    r1.font.bold = True
    r1.font.size = Pt(9.5)
    r1.font.color.rgb = RGBColor(30, 58, 138)
    r2 = p.add_run(body_text)
    r2.font.size = Pt(9.5)
    r2.font.color.rgb = RGBColor(30, 41, 59)
    doc.add_paragraph().paragraph_format.space_after = Pt(4)


def add_running_header_footer(section):
    header = section.header
    p_hdr = header.paragraphs[0]
    p_hdr.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    p_hdr.paragraph_format.space_after = Pt(0)

    # Two column header via tab stop
    p_hdr_pPr = p_hdr._p.get_or_add_pPr()
    p_hdr_tabs = parse_xml(
        f'<w:tabs {nsdecls("w")}>'
        f'<w:tab w:val="right" w:pos="9360"/>'
        f'</w:tabs>'
    )
    p_hdr_pPr.append(p_hdr_tabs)

    r_l = p_hdr.add_run("MCA Academic Assistant Chatbot using Natural Language Processing\t")
    r_l.font.name = "Times New Roman"
    r_l.font.size = Pt(8.5)
    r_l.font.color.rgb = RGBColor(100, 116, 139)

    r_r = p_hdr.add_run("Technical Report & System Documentation")
    r_r.font.name = "Times New Roman"
    r_r.font.size = Pt(8.5)
    r_r.font.color.rgb = RGBColor(100, 116, 139)

    # Footer
    footer = section.footer
    p_ftr = footer.paragraphs[0]
    p_ftr.paragraph_format.space_before = Pt(0)
    p_ftr_pPr = p_ftr._p.get_or_add_pPr()
    p_ftr_tabs = parse_xml(
        f'<w:tabs {nsdecls("w")}>'
        f'<w:tab w:val="right" w:pos="9360"/>'
        f'</w:tabs>'
    )
    p_ftr_pPr.append(p_ftr_tabs)

    r_fl = p_ftr.add_run("Academic AI Project | Next.js 16 - FastAPI - SQLite - Rule-Based NLP\t")
    r_fl.font.name = "Times New Roman"
    r_fl.font.size = Pt(8.5)
    r_fl.font.color.rgb = RGBColor(100, 116, 139)

    r_fr = p_ftr.add_run("Page ")
    r_fr.font.name = "Times New Roman"
    r_fr.font.size = Pt(8.5)
    r_fr.font.color.rgb = RGBColor(100, 116, 139)

    fld1 = parse_xml(f'<w:fldSimple {nsdecls("w")} w:instr="PAGE"/>')
    p_ftr._p.append(fld1)

    r_fr2 = p_ftr.add_run(" of ")
    r_fr2.font.name = "Times New Roman"
    r_fr2.font.size = Pt(8.5)
    r_fr2.font.color.rgb = RGBColor(100, 116, 139)

    fld2 = parse_xml(f'<w:fldSimple {nsdecls("w")} w:instr="NUMPAGES"/>')
    p_ftr._p.append(fld2)


print("Technical report building utilities initialized.")
