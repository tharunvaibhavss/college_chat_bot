"""
Generator script for MCA Academic Assistant Project Report (.docx and .pdf).
Built for PSG College of Arts & Science, Department of MCA.
Authoritative source data: Academic Calendar 2026-2027 & MCA Semester III Timetable.
"""
import os
import sys
import docx
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
from docx.oxml import parse_xml, OxmlElement
from docx.oxml.ns import nsdecls, qn
import win32com.client


def set_cell_background(cell, fill_hex):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{fill_hex}"/>')
    tcPr.append(shd)


def set_cell_margins(cell, top=120, bottom=120, left=150, right=150):
    tcPr = cell._tc.get_or_add_tcPr()
    tcMar = parse_xml(
        f'<w:tcMar {nsdecls("w")}>'
        f'<w:top w:w="{top}" w:type="dxa"/>'
        f'<w:bottom w:w="{bottom}" w:type="dxa"/>'
        f'<w:left w:w="{left}" w:type="dxa"/>'
        f'<w:right w:w="{right}" w:type="dxa"/>'
        f'</w:tcMar>'
    )
    tcPr.append(tcMar)


def add_hyperlink(paragraph, url, text, color="004B87", underline=True, bold=False):
    part = paragraph.part
    r_id = part.relate_to(url, docx.opc.constants.RELATIONSHIP_TYPE.HYPERLINK, is_external=True)

    hyperlink = parse_xml(
        f'<w:hyperlink {nsdecls("w")} r:id="{r_id}" '
        f'xmlns:r="http://schemas.openxmlformats.org/officeDocument/2006/relationships"/>'
    )
    new_run = parse_xml(f'<w:r {nsdecls("w")}/>')
    rPr = parse_xml(f'<w:rPr {nsdecls("w")}/>')

    if color:
        c = parse_xml(f'<w:color {nsdecls("w")} w:val="{color}"/>')
        rPr.append(c)
    if underline:
        u = parse_xml(f'<w:u {nsdecls("w")} w:val="single"/>')
        rPr.append(u)
    if bold:
        b = parse_xml(f'<w:b {nsdecls("w")}/>')
        rPr.append(b)

    # Set font to Times New Roman
    rFonts = parse_xml(f'<w:rFonts {nsdecls("w")} w:ascii="Times New Roman" w:hAnsi="Times New Roman"/>')
    rPr.append(rFonts)

    new_run.append(rPr)
    text_node = parse_xml(f'<w:t {nsdecls("w")}>{text}</w:t>')
    new_run.append(text_node)
    hyperlink.append(new_run)
    paragraph._p.append(hyperlink)
    return hyperlink


def style_table(table, col_widths=None, header_bg="1E3A8A", alt_bg="F8FAFC", border_c="CBD5E1"):
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    # Set borders
    tblPr = table._tbl.tblPr
    tblBorders = parse_xml(
        f'<w:tblBorders {nsdecls("w")}>'
        f'<w:top w:val="single" w:sz="6" w:space="0" w:color="{border_c}"/>'
        f'<w:bottom w:val="single" w:sz="6" w:space="0" w:color="{border_c}"/>'
        f'<w:left w:val="single" w:sz="6" w:space="0" w:color="{border_c}"/>'
        f'<w:right w:val="single" w:sz="6" w:space="0" w:color="{border_c}"/>'
        f'<w:insideH w:val="single" w:sz="4" w:space="0" w:color="{border_c}"/>'
        f'<w:insideV w:val="single" w:sz="4" w:space="0" w:color="{border_c}"/>'
        f'</w:tblBorders>'
    )
    tblPr.append(tblBorders)

    for row_idx, row in enumerate(table.rows):
        trPr = row._tr.get_or_add_trPr()
        trPr.append(parse_xml(f'<w:cantSplit {nsdecls("w")}/>'))
        if row_idx == 0:
            trPr.append(parse_xml(f'<w:tblHeader {nsdecls("w")}/>'))

        for col_idx, cell in enumerate(row.cells):
            cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
            set_cell_margins(cell, top=100, bottom=100, left=140, right=140)
            if col_widths and col_idx < len(col_widths):
                cell.width = col_widths[col_idx]

            if row_idx == 0:
                set_cell_background(cell, header_bg)
                for p in cell.paragraphs:
                    p.paragraph_format.space_before = Pt(2)
                    p.paragraph_format.space_after = Pt(2)
                    for r in p.runs:
                        r.font.bold = True
                        r.font.color.rgb = RGBColor(255, 255, 255)
                        r.font.name = "Times New Roman"
                        r.font.size = Pt(10)
            else:
                if row_idx % 2 == 1 and alt_bg:
                    set_cell_background(cell, alt_bg)
                for p in cell.paragraphs:
                    p.paragraph_format.space_before = Pt(2)
                    p.paragraph_format.space_after = Pt(2)
                    for r in p.runs:
                        r.font.name = "Times New Roman"
                        r.font.size = Pt(9.5)
                        r.font.color.rgb = RGBColor(30, 41, 59)


def add_screenshot_placeholder(doc, title, instruction):
    tbl = doc.add_table(rows=1, cols=1)
    tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    cell = tbl.cell(0, 0)
    cell.width = Inches(6.2)
    set_cell_background(cell, "F1F5F9")
    set_cell_margins(cell, top=280, bottom=280, left=200, right=200)

    tcPr = cell._tc.get_or_add_tcPr()
    tcBorders = parse_xml(
        f'<w:tcBorders {nsdecls("w")}>'
        f'<w:top w:val="dashed" w:sz="8" w:space="0" w:color="94A3B8"/>'
        f'<w:bottom w:val="dashed" w:sz="8" w:space="0" w:color="94A3B8"/>'
        f'<w:left w:val="dashed" w:sz="8" w:space="0" w:color="94A3B8"/>'
        f'<w:right w:val="dashed" w:sz="8" w:space="0" w:color="94A3B8"/>'
        f'</w:tcBorders>'
    )
    tcPr.append(tcBorders)

    p1 = cell.paragraphs[0]
    p1.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r1 = p1.add_run(instruction)
    r1.font.bold = True
    r1.font.name = "Times New Roman"
    r1.font.size = Pt(11)
    r1.font.color.rgb = RGBColor(71, 85, 105)

    p2 = cell.add_paragraph()
    p2.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r2 = p2.add_run(f"System Module: {title}")
    r2.font.italic = True
    r2.font.name = "Times New Roman"
    r2.font.size = Pt(9.5)
    r2.font.color.rgb = RGBColor(100, 116, 139)

    doc.add_paragraph()  # Spacing


def add_toc_field(paragraph):
    run = paragraph.add_run()
    fldSimple = parse_xml(
        f'<w:fldSimple {nsdecls("w")} w:instr="TOC \\o &quot;1-3&quot; \\h \\z \\u"/>'
    )
    paragraph._p.append(fldSimple)


def add_page_number_to_footer(footer):
    p = footer.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    r1 = p.add_run("PSG College of Arts & Science  |  Page ")
    r1.font.name = "Times New Roman"
    r1.font.size = Pt(9)
    r1.font.color.rgb = RGBColor(100, 116, 139)

    fld1 = parse_xml(f'<w:fldSimple {nsdecls("w")} w:instr="PAGE"/>')
    p._p.append(fld1)

    r2 = p.add_run(" of ")
    r2.font.name = "Times New Roman"
    r2.font.size = Pt(9)
    r2.font.color.rgb = RGBColor(100, 116, 139)

    fld2 = parse_xml(f'<w:fldSimple {nsdecls("w")} w:instr="NUMPAGES"/>')
    p._p.append(fld2)


print("Helper functions defined successfully.")
